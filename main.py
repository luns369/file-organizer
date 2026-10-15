# -*- coding: utf-8 -*-
"""
AI 文件整理工具 —— 分析文件并生成 Excel 总表

用法：
    python main.py          第一步：扫描 + 逐文件分析 + 生成分类方案（等待确认）
    python main.py apply    第二步：按确认后的分类方案归位，生成 Excel 总表

流程（第三档：方案先确认，再执行；第四档：学习人工修改）：
    扫描文件夹 → 提取文字内容 → DeepSeek 逐文件分析（关键词/总结）
    → DeepSeek 只给出「分类方案」（分类名 + 定义边界 + 预计包含哪些文件）
    → 在网页界面上确认/修改方案（改名、增删、合并、调整文件归属）
    → 再次运行 apply：AI 严格按确认后的方案把文件归位 → 生成 Excel 总表

特点：
  - 已分析过的文件会写入 cache.json，重复运行不再消耗 API 费用
  - 分类方案写入 scheme.json，可反复修改确认；文件库或分类要求变化时会提示重新生成
  - 总表为矩阵布局：每个分类一列，每个文件一个格子，点击格子直接打开文件
  - 你在总表里手动改过的分类会被记为「归类偏好」（cache.json 的 baseline 对比得出），
    下一轮注入提示词，AI 越用越准；这些文件归位时也永远保留你的分类
"""

import hashlib
import json
import math
import os
import sys
import time
from pathlib import Path
from urllib.parse import unquote

from openai import OpenAI
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from markitdown import MarkItDown

# ---------------- 路径与常量 ----------------
# 数据目录：源码运行时用脚本所在目录；打包成 exe 后自动改用 exe 所在目录，
# 这样用户把 exe 拷到哪里，配置/缓存/日志就存在哪里（可整体分发）。
if getattr(sys, "frozen", False):
    BASE_DIR = Path(sys.executable).resolve().parent
else:
    BASE_DIR = Path(__file__).resolve().parent
CONFIG_FILE = BASE_DIR / "config.json"
CACHE_FILE = BASE_DIR / "cache.json"
SCHEME_FILE = BASE_DIR / "scheme.json"  # 第三档：AI 生成的分类方案（等待人工确认）
GLOBAL_KEY = "__global__"  # cache.json 里存全局分类结果的固定键

SUPPORTED_EXTS = {".doc", ".docx", ".pdf", ".ppt", ".pptx", ".txt"}
MAX_TEXT_CHARS = 4000  # 发给 AI 的内容上限，控制 token 消耗
# 失败类结果不进缓存（或视为需要重试）：否则 API Key 失效一次，
# 失败结果会被永久缓存，之后每次运行都"缓存命中"直接跳过，永远无法恢复
FAIL_CATEGORIES = {"分析失败", "解析失败", "无法识别"}
MAX_EXAMPLES = 20  # 最多给 AI 看多少条「人工纠正」示例，控制 token 消耗


# ---------------- 配置与缓存 ----------------
def load_config():
    if not CONFIG_FILE.exists():
        return {}  # 首次使用：还没有配置文件，由图形界面引导填写
    try:
        return json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        sys.exit(f"config.json 格式有误：{e}\n注意路径里的反斜杠要写成两个，例如 C:\\\\Users\\\\test")


def load_cache():
    if CACHE_FILE.exists():
        return json.loads(CACHE_FILE.read_text(encoding="utf-8"))
    return {}


def save_cache(cache):
    CACHE_FILE.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")


def cache_key(path):
    """用文件内容的 MD5 作为缓存标识：重命名、移动都不影响缓存，相同内容只分析一次。"""
    try:
        h = hashlib.md5()
        with open(path, "rb") as f:
            for chunk in iter(lambda: f.read(1024 * 1024), b""):
                h.update(chunk)
        return h.hexdigest()
    except OSError:
        return legacy_key(path)  # 文件被占用读不了时退化为旧标识，后续解析会标为失败


def legacy_key(path):
    """旧版缓存标识：文件名+大小+修改时间（仅用于迁移旧缓存）。"""
    st = path.stat()
    return f"{path.name}|{st.st_size}|{int(st.st_mtime)}"


# ---------------- AI 分析 ----------------
def build_rules_block(rules):
    """把用户自定义分类要求拼成提示词段落（第一档：自定义分类要求）。
    后续档位的规则来源（场景预设细化、人工纠正示例、方案确认结果）
    都通过这一段注入提示词，结构不变。"""
    rules = (rules or "").strip()
    if not rules:
        return ""
    return f"用户额外要求（优先级最高，分类时必须遵守；与默认规则冲突时以本要求为准）：\n{rules}\n\n"


def build_examples_block(examples):
    """第四档：把用户「手动改过的归类」拼成偏好示例段落，注入提示词，越用越准。
    examples: [(文件名, 总结, 用户选择的分类, AI 原本建议的分类)]"""
    examples = list(examples or [])[:MAX_EXAMPLES]
    if not examples:
        return ""
    lines = "\n".join(f"- 「{name}」→ {cat}" for name, _sum, cat, _old in examples)
    return (f"用户以往整理时手动调整过的归类（{len(examples)} 条，这是用户的真实偏好，优先级仅次于上面的额外要求）：\n"
            f"{lines}\n"
            f"请优先沿用这些分类名和归类习惯：涉及同类内容时，分类名与上述写法保持一致；"
            f"若这些分类明显不适用于当前文件库，可以调整，但要保持相近的归类倾向。\n\n")


def build_prompt(filename, text, existing_categories, rules=""):
    cats = "、".join(sorted(existing_categories)) if existing_categories else "（暂无）"
    return f"""文件名：{filename}

文件内容（节选）：
{text}

已有分类：{cats}

{build_rules_block(rules)}请完成三件事：
1. 提取 3 个关键词；
2. 用一句话（30 字以内）总结这个文件的内容；
3. 给出一个分类名：2-6 个字的简短中文词，从内容领域或使用场景命名（如"AI学习""工作文档""烹饪食谱""医疗""合同"）。已有分类里有合适的就直接复用，都不合适再新建。

严格按以下 JSON 格式返回，不要输出其他内容：
{{"keywords": ["关键词1", "关键词2", "关键词3"], "summary": "一句话总结", "category": "分类名"}}"""


def normalize_result(data):
    keywords = data.get("keywords") or []
    if isinstance(keywords, str):
        keywords = [keywords]
    return {
        "keywords": [str(k) for k in keywords][:3],
        "summary": str(data.get("summary", "")).strip(),
        "category": str(data.get("category", "未分类")).strip() or "未分类",
    }


def analyze(client, converter, path, model, existing_categories, rules=""):
    """提取文件文字并调用 DeepSeek 分析。任何一步失败都不会中断整体流程。"""
    try:
        text = (converter.convert(str(path)).text_content or "").strip()
    except Exception as e:
        return {"keywords": [], "summary": f"文件解析失败：{e}", "category": "解析失败"}
    if not text:
        return {"keywords": [], "summary": "未能提取到文字内容", "category": "无法识别"}

    try:
        resp = client.chat.completions.create(
            model=model,
            messages=[
                {"role": "system", "content": "你是专业的文件整理助手，擅长快速理解文档内容并归类。始终输出合法 JSON。"},
                {"role": "user", "content": build_prompt(path.name, text[:MAX_TEXT_CHARS], existing_categories, rules)},
            ],
            response_format={"type": "json_object"},
            temperature=0.3,
        )
        return normalize_result(json.loads(resp.choices[0].message.content))
    except Exception as e:
        return {"keywords": [], "summary": f"AI 分析失败：{e}", "category": "分析失败"}


# ---------------- Excel 生成 ----------------
def make_file_url(path):
    """文件的 file:/// 链接：单元格挂上它，点击即可用默认程序打开文件。"""
    return "file:///" + str(path).replace("\\", "/")


def norm_path(p):
    """路径标准化，用于匹配（Windows 不区分大小写）。"""
    return os.path.normcase(os.path.normpath(str(p)))


def path_from_cell(cell):
    """从单元格的超链接还原文件路径（兼容 Excel 重新保存后的百分号编码）。"""
    if not cell.hyperlink or not cell.hyperlink.target:
        return None
    t = cell.hyperlink.target
    if t.lower().startswith("file:///"):
        t = t[8:]
    return unquote(t)


def write_excel(excel_path, rows):
    """矩阵布局：每个分类一列，每个文件一个格子；格子本身是可点击的文件链接。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "文件总表"

    groups = {}
    for r in rows:
        groups.setdefault(r["category"], []).append(r)

    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")
    wrap_top = Alignment(wrap_text=True, vertical="top")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # 文件多的分类排前面，方便查看
    categories = sorted(groups, key=lambda c: (-len(groups[c]), c))
    row_lines = {}
    for col, category in enumerate(categories, start=1):
        h = ws.cell(row=1, column=col, value=category)
        h.fill = header_fill
        h.font = header_font
        h.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = 46

        for row_i, r in enumerate(sorted(groups[category], key=lambda x: x["name"].lower()), start=2):
            text = (f"{r['name']}\n"
                    f"关键词：{'、'.join(r['keywords'])}\n"
                    f"总结：{r['summary']}\n"
                    f"修改日期：{r['mtime_str']}")
            c = ws.cell(row=row_i, column=col, value=text)
            c.alignment = wrap_top
            c.border = border
            c.hyperlink = make_file_url(r["path"])  # 点击格子直接打开文件
            # 估算这个格子需要的行数（文件名和总结会折行）
            lines = math.ceil(len(r["name"]) / 20) + math.ceil(len(r["summary"]) / 20) + 2
            row_lines[row_i] = max(row_lines.get(row_i, 0), lines)

    for row_i, lines in row_lines.items():
        ws.row_dimensions[row_i].height = lines * 14 + 8
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"
    wb.save(excel_path)


# ---------------- 分类方案：先出方案，确认后归位（第三档） ----------------
def fingerprint(keys, rules="", examples_str=""):
    """文件库指纹：文件有增删改、自定义要求或人工偏好变化时，自动重新做一次全局分类。"""
    raw = ("\n".join(sorted(keys))
           + "\n@@rules@@\n" + (rules or "").strip()
           + "\n@@examples@@\n" + (examples_str or "").strip())
    return hashlib.md5(raw.encode("utf-8")).hexdigest()


def file_key(name, summary):
    """人工偏好的稳定标识：文件名 + 一句话总结。
    不用路径，因为归档后文件会被移到子文件夹里，路径会变；总结来自缓存，归档不影响。"""
    return f"{name}\x1f{summary}"


def examples_str(examples):
    """把示例列表压成一个字符串，专用于指纹计算（排序固定，避免无谓的重做）。"""
    return "\n".join(f"{name}\x1f{cat}"
                     for name, _sum, cat, _old in sorted(examples or [], key=lambda x: x[0]))


def get_baseline():
    """上一轮 AI 归位的分类基线：{文件名|总结: AI 给的分类}。
    用户现在的总表分类与它不一致，就说明那一格是人工改的。"""
    return (load_cache().get(GLOBAL_KEY) or {}).get("baseline") or {}


def collect_corrections(excel_path, files, cache, keys):
    """第四档：找出用户在总表里手动改过的分类。
    返回 (示例列表 [(文件名, 总结, 用户分类, AI原建议)], 是否已有基线)。
    没有基线时无法判断改动，示例为空。"""
    by_path, by_namesum = load_manual_categories(excel_path)
    if not by_path and not by_namesum:
        return [], False
    baseline = get_baseline()
    if not baseline:
        return [], False  # 第一次运行本版本：先把当前总表当作基线，之后才谈得上"改动"
    out = []
    for fp in files:
        info = cache.get(keys[fp]) or {}
        summary = info.get("summary") or ""
        current = by_path.get(norm_path(fp)) or by_namesum.get((fp.name, summary))
        if not current:
            continue
        old = baseline.get(file_key(fp.name, summary))
        if old and old != current:
            out.append((fp.name, summary, current, old))
    # 手工改过的文件多的库，只取前 MAX_EXAMPLES 条，够 AI 看出偏好
    return sorted(out, key=lambda x: x[0])[:MAX_EXAMPLES], True


def init_baseline_from_excel(cache, files, file_keys, excel_path):
    """首次启用学习功能时，把现有总表里的分类记成基线（视为既成事实），
    此后你再调整总表，改动就会被识别为"人工纠正"并学习。"""
    by_path, by_namesum = load_manual_categories(excel_path)
    baseline = {}
    for fp in files:
        info = cache.get(file_keys[fp]) or {}
        summary = info.get("summary") or ""
        cat = by_path.get(norm_path(fp)) or by_namesum.get((fp.name, summary))
        if cat:
            baseline[file_key(fp.name, summary)] = cat
    if not baseline:
        return 0
    cache[GLOBAL_KEY] = dict(cache.get(GLOBAL_KEY) or {}, baseline=baseline)
    save_cache(cache)
    return len(baseline)


def set_baseline(cache, files, file_keys, mapping):
    """记录本轮 AI 归位的结果，作为下一次比对"人工改动"的基线。"""
    baseline = {}
    for i, fp in enumerate(files, 1):
        info = cache.get(file_keys[fp]) or {}
        cat = (mapping.get(str(i))
               or info.get("category")
               or "未分类")
        baseline[file_key(fp.name, info.get("summary") or "")] = cat
    cache[GLOBAL_KEY] = dict(cache.get(GLOBAL_KEY) or {}, baseline=baseline)
    return baseline


def clean_categories(raw_categories, infos):
    """规范化 AI 返回的分类列表：去空名、去重名、把文件编号翻译成文件名。"""
    out, seen = [], set()
    for c in raw_categories or []:
        if not isinstance(c, dict):
            continue
        name = str(c.get("name", "")).strip()
        if not name or name in seen:
            continue
        seen.add(name)
        names = []
        for i in c.get("files") or []:
            try:
                n = int(i)
            except (TypeError, ValueError):
                continue
            if 1 <= n <= len(infos):
                fn = infos[n - 1]["name"]
                if fn not in names:
                    names.append(fn)
        out.append({"name": name, "desc": str(c.get("desc", "")).strip(), "files": names})
    return out


def _chat_json(client, model, system, prompt):
    """统一的 JSON 模式调用。返回解析后的 dict，异常由调用方兜底。"""
    resp = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": system},
            {"role": "user", "content": prompt},
        ],
        response_format={"type": "json_object"},
        temperature=0.3,
    )
    return json.loads(resp.choices[0].message.content)


def _file_lines(infos):
    return "\n".join(
        f"{i}. {r['name']} | 关键词：{'、'.join(r['keywords']) or '无'} | 总结：{r['summary']}"
        for i, r in enumerate(infos, 1)
    )


def propose_scheme(client, model, infos, rules="", examples=None):
    """第一步：只让 AI 给出分类方案（分类名 + 定义边界 + 预计包含哪些文件）。
    返回 [{name, desc, files:[文件名]}]；失败返回 []，由调用方兜底。"""
    prompt = f"""这是一个文件夹里全部 {len(infos)} 个文件的信息：

{_file_lines(infos)}

{build_rules_block(rules)}{build_examples_block(examples)}请先只设计一套分类方案，不要急着下最终结论。要求：
1. 设计 5-12 个分类（若上方用户要求对数量、粒度、分类维度另有说明，以用户要求为准），互斥、粒度适中；
2. 分类名为 2-8 个字的简短中文词，从内容领域或使用场景命名，概括内容主题，不要照抄文件名或文件夹名；
3. 每个分类给出"定义边界"：一句话说明这一类装什么、与相邻分类怎么区分，30 字以内；
4. 给出预计归入该分类的文件编号，同一个编号不要重复出现在多个分类，尽量覆盖所有文件编号。

严格按以下 JSON 格式返回，不要输出其他内容：
{{"categories": [{{"name": "分类名", "desc": "定义边界", "files": [1, 5, 7]}}]}}"""
    try:
        data = _chat_json(
            client, model,
            "你是专业的文件整理专家，擅长为整个文件库设计统一、清晰的分类体系。始终输出合法 JSON。",
            prompt,
        )
        return clean_categories(data.get("categories") or data.get("assignments") or [], infos)
    except Exception as e:
        print(f"AI 生成分类方案失败：{e}")
        return []


def assign_by_scheme(client, model, infos, categories, rules="", examples=None):
    """第二步：按用户确认后的分类方案，把每个文件归位。
    返回 {文件编号: 分类名}；方案外的名字一律落到"未分类"。"""
    cat_lines = "\n".join(
        f"- {c['name']}：{c.get('desc') or '（无说明）'}" for c in categories
    )
    prompt = f"""这是一个文件夹里全部 {len(infos)} 个文件的信息：

{_file_lines(infos)}

用户已确认的分类方案（必须严格使用下列分类名，不得新增、改名或合并）：
{cat_lines}

{build_rules_block(rules)}{build_examples_block(examples)}请把每个文件归入上面方案中最合适的一类。内容相似的文件必须归入同一个分类；方案里确实没有合适归属的，才用"未分类"。

严格按以下 JSON 格式返回，键是文件编号（每个文件都要有）：
{{"assignments": {{"1": "分类名", "2": "分类名"}}}}"""
    try:
        data = _chat_json(
            client, model,
            "你是专业的文件整理专家，严格按用户给定的分类方案归类。始终输出合法 JSON。",
            prompt,
        )
        raw = data.get("assignments", data)
        valid = {c["name"] for c in categories}
        result = {}
        for k, v in raw.items():
            name = str(v).strip()
            result[str(k)] = name if name in valid else "未分类"
        return result
    except Exception as e:
        print(f"按方案归位失败，改用每个文件各自的建议分类。原因：{e}")
        return {}


# ---------------- 分类方案的读写 ----------------
def save_scheme(scheme):
    SCHEME_FILE.write_text(json.dumps(scheme, ensure_ascii=False, indent=2), encoding="utf-8")


def load_scheme():
    if not SCHEME_FILE.exists():
        return {}
    try:
        return json.loads(SCHEME_FILE.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return {}


def print_scheme(categories):
    print("\n========== 分类方案 ==========")
    for i, c in enumerate(categories, 1):
        names = c.get("files") or []
        print(f"{i}. {c['name']}（预计 {len(names)} 个文件）")
        if c.get("desc"):
            print(f"   定义：{c['desc']}")
    print("==============================")


def fallback_scheme(files, cache, keys):
    """AI 生成方案失败时的兜底：用逐文件分析得到的分类拼一个初始方案。"""
    groups = {}
    for fp in files:
        info = cache.get(keys[fp]) or {}
        groups.setdefault(info.get("category") or "未分类", []).append(fp.name)
    return [{"name": n, "desc": "", "files": fs}
            for n, fs in sorted(groups.items(), key=lambda x: -len(x[1]))]


def load_manual_categories(excel_path):
    """读取已有总表里人工调整过的分类，重跑时保留。
    返回两个字典：按路径匹配（新版矩阵布局）、按 文件名+总结 匹配（旧版行式布局）。"""
    by_path, by_namesum = {}, {}
    if not excel_path.exists():
        return by_path, by_namesum
    try:
        ws = load_workbook(excel_path).active
        headers = [c.value for c in ws[1]]
        if "分类" in headers and "文件名" in headers:
            # 旧版行式布局
            h = {v: i for i, v in enumerate(headers)}
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[h["文件名"]] and row[h["分类"]]:
                    key = (str(row[h["文件名"]]), str(row[h["一句话总结"]] or ""))
                    by_namesum[key] = str(row[h["分类"]])
        else:
            # 新版矩阵布局：表头即分类，带超链接的格子即文件
            for col_cells in ws.iter_cols(min_row=1):
                category = col_cells[0].value
                if not category:
                    continue
                for c in col_cells[1:]:
                    p = path_from_cell(c)
                    if p:
                        by_path[norm_path(p)] = str(category)
    except Exception as e:
        # 读不出来就中止，而不是重新生成——否则你手动改过的分类会被静默覆盖
        sys.exit(f"无法读取已有总表：{e}\n请关闭 Excel/WPS 里打开的总表后重新运行。（为保护你手动改过的分类，本次未重新生成）")
    return by_path, by_namesum


# ---------------- 主流程 ----------------

def collect_files(config):
    """读取配置、扫描待整理文件。"""
    if not config.get("source_folder"):
        sys.exit("还没有配置要整理的文件夹。\n请在图形界面「基本设置」里填写文件夹路径并保存，或直接编辑 config.json 的 source_folder 字段。")
    folder = Path(config["source_folder"])
    if not folder.is_dir():
        sys.exit(f"找不到文件夹：{folder}\n请打开 config.json，把 source_folder 改成你要整理的文件夹路径。")
    files = sorted(
        p for p in folder.rglob("*")
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTS and not p.name.startswith("~$")
    )
    if not files:
        sys.exit("该文件夹里没有找到 Word / PDF / PPT / TXT 文件。")
    return folder, files


def prepare_cache(files):
    """算出每个文件的缓存标识，并迁移旧版缓存。返回 (cache, keys)。"""
    cache = load_cache()
    keys, migrated = {}, 0
    for fp in files:
        k = cache_key(fp)
        if k not in cache:
            old = legacy_key(fp)
            if old in cache:
                cache[k] = cache[old]  # 迁移旧缓存到新标识，避免重复调 API
                migrated += 1
        keys[fp] = k
    if migrated:
        save_cache(cache)
        print(f"已迁移旧版缓存 {migrated} 条\n")
    return cache, keys


def make_client(config):
    if "在这里" in config.get("api_key", ""):
        sys.exit("请先在 config.json 里填入你的 DeepSeek API 密钥（api_key）。")
    return OpenAI(api_key=config["api_key"], base_url=config.get("base_url", "https://api.deepseek.com"))


def pending_files(files, cache, keys):
    """还没分析过（或上次失败）的文件。"""
    return [fp for fp in files
            if keys[fp] not in cache or cache[keys[fp]].get("category") in FAIL_CATEGORIES]


def run_analysis(config, files, cache, keys, client, rules):
    """逐文件 AI 分析，结果写入缓存；失败结果不落盘，下次自动重试。"""
    model = config.get("model", "deepseek-chat")
    converter = MarkItDown()
    existing_categories = {v.get("category", "") for v in cache.values()} - {""}
    total = len(files)
    for i, fp in enumerate(files, 1):
        if keys[fp] in cache and cache[keys[fp]].get("category") not in FAIL_CATEGORIES:
            print(f"[{i}/{total}] 缓存命中：{fp.name}")
            continue
        print(f"[{i}/{total}] AI 分析中：{fp.name}")
        result = analyze(client, converter, fp, model, existing_categories, rules)
        if result["category"] in FAIL_CATEGORIES:
            cache[keys[fp]] = result  # 仅内存中保留，供本次使用
            save_cache({k: v for k, v in cache.items()
                        if k == GLOBAL_KEY or v.get("category") not in FAIL_CATEGORIES})
        else:
            cache[keys[fp]] = result
            save_cache(cache)
        existing_categories.add(result["category"])


def build_infos(files, cache, keys):
    """拼出发给 AI 的文件信息列表（编号顺序与 files 一致）。"""
    infos = []
    for fp in files:
        info = cache.get(keys[fp]) or {}
        infos.append({
            "name": fp.name,
            "keywords": info.get("keywords") or [],
            "summary": info.get("summary") or "",
            "category": info.get("category") or "",
        })
    return infos


def write_final_excel(config, folder, files, cache, keys, mapping, keep_manual=None):
    """按 mapping 生成总表。
    keep_manual=None：保守模式，总表里读到的分类全部保留（无基线时的兼容行为）；
    keep_manual=集合：只有集合里的文件（= 用户在总表里手动改过的那几个）保留其分类，
    其余一律按确认后的方案归位——这样"人工修改"和"方案归位"不会互相打架。
    返回 (excel 路径, 保留人工分类的个数)。"""
    excel_path = folder / config.get("excel_name", "文件整理总表.xlsx")
    manual_path, manual_namesum = load_manual_categories(excel_path)

    rows, n_manual = [], 0
    for i, fp in enumerate(files, 1):
        info = cache.get(keys[fp]) or {}
        summary = info.get("summary", "")
        manual_cat = (manual_path.get(norm_path(fp))
                      or manual_namesum.get((fp.name, summary)))
        if keep_manual is None:
            keep = manual_cat is not None
        else:
            keep = manual_cat is not None and file_key(fp.name, summary) in keep_manual
        if keep:
            n_manual += 1
        category = ((manual_cat if keep else None)
                    or mapping.get(str(i))
                    or info.get("category")
                    or "未分类")
        rows.append({
            "category": category,
            "name": fp.name,
            "keywords": info.get("keywords") or [],
            "summary": summary,
            "mtime_str": time.strftime("%Y-%m-%d %H:%M", time.localtime(fp.stat().st_mtime)),
            "path": fp,
        })

    try:
        write_excel(excel_path, rows)
    except PermissionError:
        sys.exit("无法写入 Excel：总表正被打开，请先关闭它再重新运行。")
    return excel_path, n_manual


def do_analyze():
    """第一步：扫描 + 逐文件分析 + 生成分类方案（不生成总表，等用户确认）。"""
    config = load_config()
    folder, files = collect_files(config)
    print(f"共找到 {len(files)} 个文件\n")

    cache, keys = prepare_cache(files)
    rules = (config.get("custom_rules") or "").strip()
    excel_path = folder / config.get("excel_name", "文件整理总表.xlsx")

    # 第四档：先看你上次在总表里手动改过什么，作为分类偏好喂给 AI
    corrections, has_baseline = collect_corrections(excel_path, files, cache, keys)
    if corrections:
        print(f"已学习你手动调整过的 {len(corrections)} 条归类偏好：")
        for name, _sum, cat, old in corrections[:5]:
            print(f"  - {name} → {cat}（AI 原建议：{old}）")
        if len(corrections) > 5:
            print(f"  …… 共 {len(corrections)} 条")
        print()
    elif excel_path.exists() and not has_baseline:
        n = init_baseline_from_excel(cache, files, keys, excel_path)
        if n:
            print(f"已把现有总表里的 {n} 个分类记录为基线，你之后再改动总表就会被学习到。\n")

    fp_hash = fingerprint(list(keys.values()), rules, examples_str(corrections))
    todo = pending_files(files, cache, keys)

    scheme = load_scheme()
    scheme_fresh = bool(scheme.get("categories")) and scheme.get("fingerprint") == fp_hash

    client = None
    if todo or not scheme_fresh:
        client = make_client(config)
    if todo:
        print(f"开始逐个分析 {len(todo)} 个文件……")
        run_analysis(config, files, cache, keys, client, rules)
        print()

    if scheme_fresh:
        print("分类方案已是最新（文件库、分类要求与分类偏好都没变化），直接复用。")
        print_scheme(scheme["categories"])
    else:
        print("正在根据整个文件库设计分类方案……")
        if rules:
            print("已加载自定义分类要求")
        if corrections:
            print(f"已带入 {len(corrections)} 条你的历史归类偏好")
        model = config.get("model", "deepseek-chat")
        cats = propose_scheme(client, model, build_infos(files, cache, keys), rules, corrections)
        if not cats:
            print("已改用每个文件各自的建议分类，拼出一个初始方案，请手工调整。")
            cats = fallback_scheme(files, cache, keys)
        save_scheme({
            "fingerprint": fp_hash,
            "rules": rules,
            "created_at": time.strftime("%Y-%m-%d %H:%M:%S"),
            "categories": cats,
            # 界面用：告诉用户这次参考了多少条他手动改过的归类
            "learned": [{"file": n, "category": c, "was": o} for n, _s, c, o in corrections],
        })
        print_scheme(cats)

    print("\n下一步：在页面「分类方案」里检查并修改（改名 / 增删 / 合并 / 移动文件），")
    print("确认后点「按方案生成总表」。")


def do_apply():
    """第二步：按确认后的方案归位，生成 Excel 总表。"""
    config = load_config()
    folder, files = collect_files(config)

    scheme = load_scheme()
    cats = [c for c in (scheme.get("categories") or []) if str(c.get("name", "")).strip()]
    if not cats:
        sys.exit("还没有分类方案。请先点「开始分析」生成方案，确认后再点「按方案生成总表」。")

    cache, keys = prepare_cache(files)
    rules = (config.get("custom_rules") or "").strip()
    excel_path = folder / config.get("excel_name", "文件整理总表.xlsx")

    corrections, has_baseline = collect_corrections(excel_path, files, cache, keys)
    fp_hash = fingerprint(list(keys.values()), rules, examples_str(corrections))
    if scheme.get("fingerprint") != fp_hash:
        print("提示：文件库、分类要求或人工偏好与生成方案时不同，将按当前文件库和这份方案归位。\n")

    client = make_client(config)
    todo = pending_files(files, cache, keys)
    if todo:
        print(f"有 {len(todo)} 个文件还没分析过，先补上……")
        run_analysis(config, files, cache, keys, client, rules)
        print()

    if corrections:
        print(f"已带上你手动调整过的 {len(corrections)} 条归类偏好辅助归位。")

    print(f"按确认后的方案归位（共 {len(cats)} 个分类）：")
    for c in cats:
        print(f"  - {c['name']}")
    model = config.get("model", "deepseek-chat")
    mapping = assign_by_scheme(client, model, build_infos(files, cache, keys), cats, rules, corrections)
    if not mapping:
        print("AI 归位失败，已改用每个文件各自的建议分类兜底。")

    set_baseline(cache, files, keys, mapping)   # 记录本轮 AI 结果，供下次识别"你改了哪几个"
    cache[GLOBAL_KEY] = dict(cache[GLOBAL_KEY], fingerprint=fp_hash, mapping=mapping, scheme=cats)
    save_cache(cache)

    # 你手动改过的文件保留你的分类，其余严格按方案归位
    keep_manual = ({file_key(n, s) for n, s, _c, _o in corrections}
                   if has_baseline else None)
    excel_path, n_manual = write_final_excel(config, folder, files, cache, keys, mapping, keep_manual)
    print(f"\n完成！总表已生成：{excel_path}")
    print(f"分类来源：方案归位 {len(files) - n_manual} 个，保留你手动修改的 {n_manual} 个")
    print("接下来：打开总表检查并修改分类，然后点「开始归档」。")


def main(mode="analyze"):
    if mode == "apply":
        do_apply()
    else:
        do_analyze()


if __name__ == "__main__":
    try:
        main(sys.argv[1] if len(sys.argv) > 1 else "analyze")
    finally:
        # 仅交互式终端运行时提示按回车（网页界面调用时会设置 FO_SILENT 自动跳过）
        if not os.environ.get("FO_SILENT") and sys.stdin.isatty():
            input("\n按回车键退出...")
