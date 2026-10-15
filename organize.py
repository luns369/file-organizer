# -*- coding: utf-8 -*-
"""
AI 文件整理工具 —— 第二步：按 Excel 总表归档文件

用法：
    python organize.py        归档：按总表「分类」列把文件移动到对应子文件夹
    python organize.py undo   撤销：把上一次归档移动的文件全部放回原位

特点：
  - 移动（不是复制），归档后自动更新格子的文件链接，点击不会失效
  - 分类名支持两级「一级/二级」（如"学习/考公"）：自动创建两层文件夹，最多两层
  - 每次归档都会写 undo.json 日志，可随时一键撤销（含清理空文件夹）
  - 重新归类：把文件格子剪切粘贴到另一列（或修改列表头），重跑本脚本即可
"""

import json
import os
import re
import shutil
import sys
from pathlib import Path
from urllib.parse import unquote

from openpyxl import load_workbook

# 数据目录：源码运行时用脚本所在目录；打包成 exe 后自动改用 exe 所在目录
if getattr(sys, "frozen", False):
    BASE_DIR = Path(sys.executable).resolve().parent
else:
    BASE_DIR = Path(__file__).resolve().parent
CONFIG_FILE = BASE_DIR / "config.json"
UNDO_FILE = BASE_DIR / "undo.json"

MAX_LEVELS = 2      # 分类层级上限，与 main.py 保持一致
DATA_START_ROW = 3  # 双行表头（一级 + 二级）时的数据起始行


def sanitize(name):
    """去掉 Windows 文件夹名不允许的字符。"""
    name = re.sub(r'[\\/:*?"<>|]', "_", str(name)).strip().strip(".")
    return name or "未分类"


def sanitize_rel(name, max_levels=MAX_LEVELS):
    """把分类名变成相对目录路径：「学习/考公」→ 学习/考公（最多两级）。
    每段单独清理非法字符，空段丢弃，防止路径穿越或过深嵌套。"""
    raw = str(name or "")
    for ch in ("\\", "／", "｜", "|", ">", "〉"):
        raw = raw.replace(ch, "/")
    parts = [sanitize(p) for p in raw.split("/") if p and p.strip()]
    if not parts:
        return Path("未分类")
    return Path(*parts[:max_levels])


def unique_path(target):
    """目标位置已有同名文件时，自动加序号：xxx_1.pdf、xxx_2.pdf……"""
    if not target.exists():
        return target
    n = 1
    while True:
        candidate = target.parent / f"{target.stem}_{n}{target.suffix}"
        if not candidate.exists():
            return candidate
        n += 1


def make_file_url(path):
    """文件的 file:/// 链接：单元格挂上它，点击即可用默认程序打开文件。"""
    return "file:///" + str(path).replace("\\", "/")


def path_from_cell(cell):
    """从单元格的超链接还原文件路径（兼容 Excel 重新保存后的百分号编码）。"""
    if not cell.hyperlink or not cell.hyperlink.target:
        return None
    t = cell.hyperlink.target
    if t.lower().startswith("file:///"):
        t = t[8:]
    return unquote(t)


def load_config():
    if not CONFIG_FILE.exists():
        return {}  # 首次使用：还没有配置文件，由图形界面引导填写
    return json.loads(CONFIG_FILE.read_text(encoding="utf-8"))


def do_organize():
    config = load_config()
    if not config.get("source_folder"):
        sys.exit("还没有配置要整理的文件夹。\n请在图形界面「基本设置」里填写文件夹路径并保存。")
    folder = Path(config["source_folder"])
    excel_path = folder / config.get("excel_name", "文件整理总表.xlsx")
    if not excel_path.exists():
        sys.exit(f"找不到总表：{excel_path}\n请先运行  python main.py  生成总表。")

    # 先确认总表可写，避免文件移动完才发现它被 Excel 占用
    try:
        with open(excel_path, "r+b"):
            pass
    except PermissionError:
        sys.exit("总表正被 Excel 打开，请先关闭它再运行本脚本。")

    wb = load_workbook(excel_path)
    ws = wb.active

    if "分类" in [c.value for c in ws[1]]:
        sys.exit("总表是旧版行式布局，请先运行  python main.py  重新生成新版总表（分类为列、每个文件一个格子）。")

    # 判断表头是几行：第 2 行就有文件链接 = 旧版单行表头；否则为两级表头（第 1 行一级、第 2 行二级）
    data_start = 2 if (ws.max_row >= 2 and any(path_from_cell(c) for c in ws[2])) else DATA_START_ROW

    # 矩阵布局：表头即分类，下面每个带超链接的格子是一个文件
    undo_log = []
    top = ""
    for col_cells in ws.iter_cols(min_row=1):
        if data_start == 2:
            header = col_cells[0].value
            cells = col_cells[1:]
        else:
            v1 = str(col_cells[0].value or "").strip()
            if v1:
                top = v1   # 一级表头是合并单元格，只有该级第一列有值，向后沿用
            sub = str(col_cells[1].value or "").strip() if len(col_cells) > 1 else ""
            header = f"{top}/{sub}" if sub else top
            cells = col_cells[2:]
        if not str(header or "").strip():
            continue
        rel = sanitize_rel(header)          # 「学习/考公」→ 学习/考公
        label = "/".join(rel.parts)
        for cell in cells:
            p = path_from_cell(cell)
            if not p:
                continue
            src = Path(p)
            if not src.exists():
                print(f"跳过（文件不存在）：{src}")
                continue
            target_dir = folder / rel
            if src.parent == target_dir:
                continue  # 已经在正确的文件夹里
            target_dir.mkdir(parents=True, exist_ok=True)
            dst = unique_path(target_dir / src.name)
            shutil.move(str(src), str(dst))
            undo_log.append({"from": str(dst), "to": str(src)})
            cell.hyperlink = make_file_url(dst)
            print(f"{src.name}  ->  {label}/")

    try:
        wb.save(excel_path)
    except PermissionError:
        sys.exit("\n文件已移动，但总表正被打开无法更新链接。请关闭 Excel 后重新运行本脚本。")

    UNDO_FILE.write_text(json.dumps(undo_log, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\n归档完成！共移动 {len(undo_log)} 个文件到 {len({Path(i['from']).parent for i in undo_log})} 个分类文件夹。")
    print("如需撤销：python organize.py undo")


def do_undo():
    if not UNDO_FILE.exists():
        sys.exit("没有可撤销的归档记录。")
    log = json.loads(UNDO_FILE.read_text(encoding="utf-8"))
    for item in reversed(log):
        src, dst = Path(item["from"]), Path(item["to"])
        if src.exists():
            dst.parent.mkdir(parents=True, exist_ok=True)
            shutil.move(str(src), str(unique_path(dst)))
            print(f"已恢复：{src.name}")
    # 清理撤销后留下的空文件夹：两级目录要先删二级、再删一级
    try:
        folder = Path(load_config().get("source_folder", "") or "")
        folder_key = os.path.normcase(str(folder))
    except Exception:
        folder, folder_key = None, None
    dirs = set()
    for i in log:
        d = Path(i["from"]).parent
        dirs.add(d)
        if folder_key:
            for anc in d.parents:          # 一路向上收集祖先目录，直到整理文件夹本身
                if os.path.normcase(str(anc)) == folder_key:
                    break
                dirs.add(anc)
    for d in sorted(dirs, key=lambda p: len(p.parts), reverse=True):
        try:
            d.rmdir()                      # 非空会失败，忽略即可
        except OSError:
            pass
    UNDO_FILE.unlink()
    print("\n撤销完成！注意：总表里的路径还是归档后的，重新运行  python main.py  即可刷新（已分析过的文件走缓存，不花 API 费用）。")


if __name__ == "__main__":
    try:
        if len(sys.argv) > 1 and sys.argv[1] == "undo":
            do_undo()
        else:
            do_organize()
    finally:
        # 仅交互式终端运行时提示按回车（网页界面调用时会设置 FO_SILENT 自动跳过）
        if not os.environ.get("FO_SILENT") and sys.stdin.isatty():
            input("\n按回车键退出...")
